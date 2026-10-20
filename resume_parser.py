"""Resume text extraction for PDF, DOCX, and image files."""
from io import BytesIO
from pathlib import Path

def ocr_available():
    try:
        import pytesseract
        from PIL import Image
        return True
    except Exception:
        return False

def extract_text(uploaded_file, filename=None):
    name=(filename or getattr(uploaded_file, 'name', '') or '').lower()
    data=uploaded_file.getvalue() if hasattr(uploaded_file,'getvalue') else uploaded_file.read()
    if name.endswith('.pdf'):
        try:
            import fitz
            doc=fitz.open(stream=data, filetype='pdf')
            if doc.needs_pass:
                raise ValueError('This PDF is password-protected. Upload an unlocked copy.')
            text='\n'.join(page.get_text('text') for page in doc).strip()
            if text:
                return text
            if ocr_available():
                import pytesseract
                from PIL import Image
                pages=[]
                for page_no in range(min(3, doc.page_count)):
                    pix=doc[page_no].get_pixmap(matrix=fitz.Matrix(2,2), alpha=False)
                    img=Image.open(BytesIO(pix.tobytes('png')))
                    pages.append(pytesseract.image_to_string(img))
                return '\n'.join(pages).strip()
            raise ValueError('This PDF appears to be scanned. Install Tesseract OCR to read scanned PDFs.')
        except ImportError:
            try:
                from pypdf import PdfReader
                reader=PdfReader(BytesIO(data))
                if reader.is_encrypted:
                    raise ValueError('This PDF is password-protected. Upload an unlocked copy.')
                return '\n'.join(page.extract_text() or '' for page in reader.pages).strip()
            except ImportError as exc:
                raise ValueError('Install PyMuPDF or pypdf to read PDF files.') from exc
    if name.endswith('.docx'):
        try:
            from docx import Document
            doc=Document(BytesIO(data))
            parts=[p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(' | '.join(cell.text for cell in row.cells))
            return '\n'.join(parts).strip()
        except Exception as exc:
            raise ValueError('Could not read this DOCX file. Try opening and resaving it in Word.') from exc
    if name.endswith(('.png','.jpg','.jpeg')):
        if not ocr_available():
            raise ValueError('Image OCR is unavailable. Install Pillow and pytesseract, plus the Tesseract application.')
        import pytesseract
        from PIL import Image, ImageOps
        try:
            img=ImageOps.exif_transpose(Image.open(BytesIO(data))).convert('RGB')
            return pytesseract.image_to_string(img).strip()
        except Exception as exc:
            raise ValueError('Could not read this image. Upload a clearer PNG or JPG.') from exc
    raise ValueError('Unsupported file type. Upload PDF, DOCX, PNG, or JPG.')
